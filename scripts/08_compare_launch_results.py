#!/usr/bin/env python3
"""Compare launch-test evidence directories and generate an XLSX report."""

import argparse
import csv
import os
import statistics
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl is required: pip3 install openpyxl", file=sys.stderr)
    sys.exit(2)


def read_csv_map(path, key):
    with open(path, newline="", encoding="utf-8") as f:
        return {row[key]: row for row in csv.DictReader(f)}


def read_csv_rows(path):
    if not path or not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_launch(path):
    rows = read_csv_map(path, "package")
    for row in rows.values():
        for key in ("t1", "t2", "t3", "t4", "t5", "avg"):
            if key in row:
                row[key] = int(row[key]) if row.get(key) else None
    return rows


def find_single(directory, rel_dir, prefix, suffix):
    path = os.path.join(directory, rel_dir)
    matches = [
        os.path.join(path, name)
        for name in os.listdir(path)
        if name.startswith(prefix) and name.endswith(suffix)
    ]
    if prefix == "apps_launch_":
        matches = [
            name for name in matches
            if "apps_launch_attempts_" not in os.path.basename(name)
            and "apps_launch_keyword_summary_" not in os.path.basename(name)
        ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one {prefix}*{suffix} under {path}, found {len(matches)}")
    return matches[0]


def mean(values):
    return round(sum(values) / len(values), 1) if values else None


def median(values):
    return round(statistics.median(values), 1) if values else None


def warm_avg(row):
    vals = [row.get(k) for k in ("t2", "t3", "t4", "t5") if row.get(k) is not None]
    return round(sum(vals) / len(vals), 1) if vals else None


def dex_state(device, pkg):
    row = device["dex"].get(pkg, {})
    parts = [row.get("filter", ""), row.get("reason", "")]
    return "/".join([p for p in parts if p])


def version_state(device, pkg):
    row = device["apk"].get(pkg, {})
    parts = [row.get("versionCode", ""), row.get("versionName", "")]
    return "/".join([p for p in parts if p])


def sha_state(device, pkg):
    return device["sha"].get(pkg, {}).get("sha256", "")


def first_nonempty(*values):
    for value in values:
        if value:
            return value
    return ""


def to_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def to_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def cv_pct(values):
    """变异系数 % = 样本标准差 / 均值 * 100。少于 2 个样本返回 None。"""
    vals = [v for v in values if v is not None]
    if len(vals) < 2:
        return None
    m = sum(vals) / len(vals)
    if not m:
        return None
    return round(statistics.stdev(vals) * 100 / m, 1)


def read_launch_window(directory):
    """读取 launch_window.txt 的 started_at/ended_at（host epoch 秒）。"""
    path = os.path.join(directory, "launch_window.txt")
    started = ended = None
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line.startswith("started_at="):
                    started = to_int(line.split("=", 1)[1])
                elif line.startswith("ended_at="):
                    ended = to_int(line.split("=", 1)[1])
    return started, ended


def summarize_timeline(directory):
    """汇总该设备 launch 窗口内的系统状态；无 timeline 返回 None。"""
    rows = read_csv_rows(os.path.join(directory, "timeline", "timeline.csv"))
    if not rows:
        return None
    started, ended = read_launch_window(directory)
    if started and ended:
        windowed = [
            r for r in rows
            if (to_int(r.get("epoch")) or 0) >= started and (to_int(r.get("epoch")) or 0) <= ended
        ]
        rows = windowed or rows  # 窗口内无采样点时退回整段，避免空表
        window_label = f"{started}->{ended}"
    else:
        window_label = "full"

    govs, cur_freqs, iowaits, thermals, dex2oat = set(), [], [], [], []
    for r in rows:
        for key, val in r.items():
            if key.endswith("_governor") and val:
                govs.add(val)
            elif key.endswith("_cur_freq"):
                iv = to_int(val)
                if iv is not None:
                    cur_freqs.append(iv)
        iw = to_float(r.get("iowait_pct"))
        if iw is not None:
            iowaits.append(iw)
        tv = to_int(r.get("thermal_max_temp"))
        if tv is not None:
            thermals.append(tv)
        dv = to_int(r.get("dex2oat_count"))
        if dv is not None:
            dex2oat.append(dv)
    return {
        "samples": len(rows),
        "window": window_label,
        "governors": ",".join(sorted(govs)),
        "cur_freq_max": max(cur_freqs) if cur_freqs else None,
        "iowait_mean": round(sum(iowaits) / len(iowaits), 1) if iowaits else None,
        "iowait_max": round(max(iowaits), 1) if iowaits else None,
        "thermal_max": max(thermals) if thermals else None,
        "dex2oat_max": max(dex2oat) if dex2oat else None,
    }


def find_optional_single(directory, rel_dir, prefix, suffix):
    path = os.path.join(directory, rel_dir)
    if not os.path.isdir(path):
        return None
    matches = [
        os.path.join(path, name)
        for name in os.listdir(path)
        if name.startswith(prefix) and name.endswith(suffix)
    ]
    return matches[0] if len(matches) == 1 else None


def build_strict_attempts(directory):
    path = find_optional_single(directory, "launch_raw", "included_strict_", ".csv")
    rows = read_csv_rows(path)
    grouped = {}
    for row in rows:
        pkg = row.get("package", "")
        if not pkg:
            continue
        grouped.setdefault(pkg, []).append(row)
    return grouped


def strict_pkg_summary(device, pkg):
    rows = device["strict"].get(pkg, [])
    totals = [v for v in (to_int(r.get("total_time_ms")) for r in rows) if v is not None and v > 0]
    displayed = [v for v in (to_int(r.get("displayed_time_ms")) for r in rows) if v is not None and v > 0]

    def avg_field(field):
        vals = [v for v in (to_int(r.get(field)) for r in rows) if v is not None]
        return mean(vals)

    first = rows[0] if rows else {}
    return {
        "strict_count": len(rows),
        "total_avg_ms": mean(totals),
        "displayed_avg_ms": mean(displayed),
        "activity": first.get("activity", ""),
        "component": first.get("component", ""),
        "displayed_activity": first.get("displayed_activity", ""),
        "final_start_activity": first.get("final_start_activity", ""),
        "version": first_nonempty(
            "/".join([p for p in (first.get("versionCode", ""), first.get("versionName", "")) if p]),
            version_state(device, pkg),
        ),
        "sha": first_nonempty(first.get("apk_sha256", ""), sha_state(device, pkg)),
        "dexopt": first_nonempty(
            "/".join([p for p in (first.get("dexopt_filter", ""), first.get("dexopt_reason", "")) if p]),
            dex_state(device, pkg),
        ),
        "keyword_target_pid_avg": avg_field("keyword_target_pid"),
        "bytehook_target_pid_avg": avg_field("bytehook_target_pid"),
        "rmonitor_target_pid_avg": avg_field("rmonitor_target_pid"),
        "shadowhook_target_pid_avg": avg_field("shadowhook_target_pid"),
        "bugly_target_pid_avg": avg_field("bugly_target_pid"),
        "total_cv_pct": cv_pct(totals),
        "total_min_ms": min(totals) if totals else None,
        "total_max_ms": max(totals) if totals else None,
        "cold_window_ms_avg": avg_field("cold_window_ms"),
        "hook_hits_in_window_avg": avg_field("hook_hits_in_window"),
        "hook_main_hits_in_window_avg": avg_field("hook_main_hits_in_window"),
        "hook_worker_hits_in_window_avg": avg_field("hook_worker_hits_in_window"),
        "hook_main_span_ms_avg": avg_field("hook_main_span_ms"),
        "metadata_note": ";".join(sorted({r.get("metadata_note", "") for r in rows if r.get("metadata_note", "")})),
    }


def values_aligned(values, require_nonempty=True):
    vals = [v for v in values if v]
    if require_nonempty and len(vals) != len(values):
        return False
    if not vals:
        return not require_nonempty
    return len(set(vals)) == 1


def add_rows(ws, rows):
    for row in rows:
        ws.append(row)


def style_workbook(wb):
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            max_len = max(len("" if cell.value is None else str(cell.value)) for cell in col)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 55)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top")


def build_device(label, directory):
    launch = read_launch(find_single(directory, "launch_raw", "apps_launch_", ".csv"))
    device = {
        "label": label,
        "dir": directory,
        "launch": launch,
        "dex": read_csv_map(os.path.join(directory, "dexopt_before.csv"), "pkg"),
        "apk": read_csv_map(os.path.join(directory, "apk_versions_before.csv"), "pkg"),
        "strict": build_strict_attempts(directory),
        "hook": build_detail_hook(directory),
        "sha": {},
    }
    sha_path = os.path.join(directory, "apk_sha256_before.csv")
    if os.path.exists(sha_path):
        device["sha"] = read_csv_map(sha_path, "pkg")
    return device


# ---------- 通用辅助 ----------
def dex_filter(device, pkg):
    """仅 dexopt 编译级别(filter)，不含 reason。"""
    return (device["dex"].get(pkg, {}) or {}).get("filter", "").strip()


def wide_cv(launch_row):
    """从宽表 t1..t5 计算 CV%。"""
    vals = [launch_row.get(k) for k in ("t1", "t2", "t3", "t4", "t5") if launch_row.get(k) is not None]
    return cv_pct(vals)


def build_detail_hook(directory):
    """从 attempts detail CSV 聚合每个 app 的主线程 hook 命中/时长(均值)，覆盖全部 app(不限 strict)。"""
    path = find_optional_single(directory, "launch_raw", "apps_launch_attempts_detail_", ".csv")
    rows = read_csv_rows(path)
    grouped = {}
    for row in rows:
        pkg = row.get("package")
        if pkg:
            grouped.setdefault(pkg, []).append(row)
    out = {}
    for pkg, rs in grouped.items():
        def avg(field):
            vals = [v for v in (to_int(r.get(field)) for r in rs) if v is not None]
            return mean(vals)
        out[pkg] = {
            "main": avg("hook_main_hits_in_window"),
            "span": avg("hook_main_span_ms"),
        }
    return out


def compute_strict_common(devices):
    """各设备都 strict 且 启动入口/Displayed/dexopt/version/sha 全对齐的 app 集合。
    返回 (pkg 列表, {pkg: [各设备 strict_pkg_summary]})。"""
    pkgs = set(devices[0]["strict"])
    for device in devices[1:]:
        pkgs &= set(device["strict"])
    common, summaries_by_pkg = [], {}
    for pkg in sorted(pkgs):
        summaries = [strict_pkg_summary(device, pkg) for device in devices]
        checks = [[s[k] for s in summaries] for k in
                  ("activity", "component", "displayed_activity", "dexopt", "version", "sha")]
        if not all(values_aligned(values) for values in checks):
            continue
        if any(s["total_avg_ms"] is None or s["displayed_avg_ms"] is None for s in summaries):
            continue
        common.append(pkg)
        summaries_by_pkg[pkg] = summaries
    return common, summaries_by_pkg


# ---------- Sheet 1: overview（全共同 app 总览，第一个设备为基准）----------
def make_overview(wb, devices, strict_set):
    ws = wb.active
    ws.title = "overview"
    base = devices[0]
    labels = [d["label"] for d in devices]
    header = ["package"]
    for label in labels:
        header += [f"{label}_首启t1", f"{label}_5次均值", f"{label}_CV%",
                   f"{label}_dexopt", f"{label}_主线程hook", f"{label}_hookSpan"]
    for label in labels[1:]:
        header += [f"{label}-{labels[0]}_Δms", f"{label}比{labels[0]}_%"]
    header += ["版本一致", "严格可比"]
    ws.append(header)

    pkgs = set(base["launch"])
    for device in devices[1:]:
        pkgs &= set(device["launch"])
    common = sorted(pkgs)
    for pkg in common:
        row = [pkg]
        for device in devices:
            lr = device["launch"].get(pkg, {})
            hk = device["hook"].get(pkg, {})
            row += [lr.get("t1"), lr.get("avg"), wide_cv(lr),
                    dex_filter(device, pkg), hk.get("main"), hk.get("span")]
        base_avg = base["launch"].get(pkg, {}).get("avg")
        for device in devices[1:]:
            dev_avg = device["launch"].get(pkg, {}).get("avg")
            if base_avg and dev_avg is not None:
                row += [dev_avg - base_avg, round((dev_avg - base_avg) * 100 / base_avg, 1)]
            else:
                row += ["", ""]
        vers = {version_state(d, pkg) for d in devices if version_state(d, pkg)}
        row += ["是" if len(vers) == 1 else "否", "是" if pkg in strict_set else "否"]
        ws.append(row)
    return common


# ---------- Sheet 2: strict（严格可比共同 app，主结论口径 + 详细列）----------
def make_strict(wb, devices, strict_common, summaries_by_pkg):
    ws = wb.create_sheet("strict")
    labels = [d["label"] for d in devices]
    header = ["package", "component", "displayed_activity", "dexopt", "version", "apk_sha256"]
    for label in labels:
        header += [
            f"{label}_strict数", f"{label}_5次均值", f"{label}_CV%",
            f"{label}_min", f"{label}_max",
            f"{label}_首启t1", f"{label}_温启t2_t5",
            f"{label}_displayed均值", f"{label}_cold_window",
            f"{label}_主线程hook", f"{label}_子线程hook", f"{label}_hookSpan",
            f"{label}_keyword命中", f"{label}_bytehook", f"{label}_rmonitor",
            f"{label}_shadowhook", f"{label}_bugly", f"{label}_note",
        ]
    for label in labels[1:]:
        header += [f"{label}-{labels[0]}_Δms", f"{label}比{labels[0]}_%", f"{label}-{labels[0]}_Δdisplayed"]
    ws.append(header)

    base = devices[0]
    for pkg in strict_common:
        summaries = summaries_by_pkg[pkg]
        s0 = summaries[0]
        row = [pkg, s0["component"], s0["displayed_activity"], s0["dexopt"], s0["version"], s0["sha"]]
        for device, s in zip(devices, summaries):
            lr = device["launch"].get(pkg, {})
            row += [
                s["strict_count"], s["total_avg_ms"], s["total_cv_pct"],
                s["total_min_ms"], s["total_max_ms"],
                lr.get("t1"), warm_avg(lr),
                s["displayed_avg_ms"], s["cold_window_ms_avg"],
                s["hook_main_hits_in_window_avg"], s["hook_worker_hits_in_window_avg"], s["hook_main_span_ms_avg"],
                s["keyword_target_pid_avg"], s["bytehook_target_pid_avg"], s["rmonitor_target_pid_avg"],
                s["shadowhook_target_pid_avg"], s["bugly_target_pid_avg"], s["metadata_note"],
            ]
        base_total = summaries[0]["total_avg_ms"]
        for s in summaries[1:]:
            total_delta = s["total_avg_ms"] - base_total
            row += [
                total_delta,
                round(total_delta * 100 / base_total, 1) if base_total else None,
                s["displayed_avg_ms"] - summaries[0]["displayed_avg_ms"],
            ]
        ws.append(row)


# ---------- Sheet 3: system_state（系统状态对齐）----------
def make_system_state(wb, devices):
    ws = wb.create_sheet("system_state")
    ws.append([
        "device", "timeline_samples", "window_epoch", "governors",
        "cur_freq_max", "iowait_mean_pct", "iowait_max_pct", "thermal_max", "dex2oat_max",
    ])
    valid = []
    for device in devices:
        s = summarize_timeline(device["dir"])
        if s is None:
            ws.append([device["label"], "timeline missing", "", "", "", "", "", "", ""])
            continue
        valid.append(s)
        ws.append([
            device["label"], s["samples"], s["window"], s["governors"],
            s["cur_freq_max"], s["iowait_mean"], s["iowait_max"], s["thermal_max"], s["dex2oat_max"],
        ])
    if len(valid) >= 2:
        ws.append([])
        ws.append(["governor_aligned_all", len({s["governors"] for s in valid}) == 1])
        ws.append([
            "note",
            "若各设备 governor 一致且 iowait/thermal/dex2oat 接近，可排除降频/限温/后台编译导致的设备级差异。",
        ])


# ---------- Sheet 4: summary（汇总，相对基准）----------
def make_summary(wb, devices, overview_common, strict_common, summaries_by_pkg):
    ws = wb.create_sheet("summary")
    labels = [d["label"] for d in devices]
    base = devices[0]
    ws.append(["metric", "value"])
    ws.append(["baseline(基准设备)", labels[0]])
    ws.append(["device_order", " -> ".join(labels)])
    ws.append(["overview_common_count", len(overview_common)])
    ws.append(["strict_common_count", len(strict_common)])
    ws.append([])

    def avg_metric(device, metric, pkgs):
        vals = [device["launch"][pkg][metric] for pkg in pkgs
                if device["launch"].get(pkg, {}).get(metric) is not None]
        return mean(vals)

    ws.append(["== overview 全共同 app（宽口径 5次均值/首启）=="])
    for device in devices:
        ws.append([f"{device['label']}_5次均值_mean", avg_metric(device, "avg", overview_common)])
        ws.append([f"{device['label']}_首启t1_mean", avg_metric(device, "t1", overview_common)])
    for device in devices[1:]:
        deltas = [device["launch"][pkg]["avg"] - base["launch"][pkg]["avg"] for pkg in overview_common
                  if base["launch"].get(pkg, {}).get("avg") and device["launch"].get(pkg, {}).get("avg") is not None]
        ws.append([f"{device['label']}-{labels[0]}_Δ均值_mean", mean(deltas)])
        ws.append([f"{device['label']}-{labels[0]}_Δ均值_median", median(deltas)])
    ws.append([])

    if strict_common:
        ws.append([f"== strict 严格可比 app（主结论口径, n={len(strict_common)}）=="])
        base_vals = [summaries_by_pkg[pkg][0]["total_avg_ms"] for pkg in strict_common]
        for i, device in enumerate(devices):
            vals = [summaries_by_pkg[pkg][i]["total_avg_ms"] for pkg in strict_common]
            ws.append([f"strict_{device['label']}_5次均值_mean", mean(vals)])
        for i, device in enumerate(devices[1:], start=1):
            deltas = [summaries_by_pkg[pkg][i]["total_avg_ms"] - base_vals[j]
                      for j, pkg in enumerate(strict_common)]
            ws.append([f"strict_{device['label']}-{labels[0]}_Δ均值_mean", mean(deltas)])
            ws.append([f"strict_{device['label']}-{labels[0]}_Δ均值_median", median(deltas)])


def main():
    parser = argparse.ArgumentParser(
        description="对比多台设备冷启动证据目录，输出标准 4-sheet XLSX (overview/strict/system_state/summary)。"
                    "第一个 --device 为基准，所有差值相对基准。")
    parser.add_argument(
        "--device", action="append", required=True,
        help="LABEL=/path/to/evidence_dir，可给 2 台或更多；第一个为基准。")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    devices = []
    for item in args.device:
        if "=" not in item:
            parser.error("--device must be LABEL=/path")
        label, directory = item.split("=", 1)
        devices.append(build_device(label, directory))
    if len(devices) < 2:
        parser.error("至少提供 2 台 --device")

    strict_common, summaries_by_pkg = compute_strict_common(devices)
    strict_set = set(strict_common)

    wb = Workbook()
    overview_common = make_overview(wb, devices, strict_set)
    make_strict(wb, devices, strict_common, summaries_by_pkg)
    make_system_state(wb, devices)
    make_summary(wb, devices, overview_common, strict_common, summaries_by_pkg)
    style_workbook(wb)
    wb.save(args.out)
    print(args.out)


if __name__ == "__main__":
    main()
