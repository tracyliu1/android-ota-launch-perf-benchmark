#!/usr/bin/env python3
"""Compare launch-test evidence directories and generate an XLSX report."""

import argparse
import csv
import os
import statistics
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl is required: pip3 install openpyxl", file=sys.stderr)
    sys.exit(2)


def read_csv_map(path, key):
    with open(path, newline="", encoding="utf-8") as f:
        return {row[key]: row for row in csv.DictReader(f)}


def read_launch(path):
    rows = read_csv_map(path, "package")
    for row in rows.values():
        for key in ("t1", "t2", "t3", "t4", "t5", "avg"):
            if key in row:
                row[key] = int(row[key]) if row.get(key) else None
    return rows


def find_single(directory, rel_dir, prefix, suffix):
    path = os.path.join(directory, rel_dir)
    matches = [
        os.path.join(path, name)
        for name in os.listdir(path)
        if name.startswith(prefix) and name.endswith(suffix)
    ]
    if prefix == "apps_launch_":
        matches = [name for name in matches if "apps_launch_attempts_" not in os.path.basename(name)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one {prefix}*{suffix} under {path}, found {len(matches)}")
    return matches[0]


def mean(values):
    return round(sum(values) / len(values), 1) if values else None


def median(values):
    return round(statistics.median(values), 1) if values else None


def warm_avg(row):
    vals = [row.get(k) for k in ("t2", "t3", "t4", "t5") if row.get(k) is not None]
    return round(sum(vals) / len(vals), 1) if vals else None


def dex_state(device, pkg):
    row = device["dex"].get(pkg, {})
    parts = [row.get("filter", ""), row.get("reason", "")]
    return "/".join([p for p in parts if p])


def version_state(device, pkg):
    row = device["apk"].get(pkg, {})
    parts = [row.get("versionCode", ""), row.get("versionName", "")]
    return "/".join([p for p in parts if p])


def sha_state(device, pkg):
    return device["sha"].get(pkg, {}).get("sha256", "")


def add_rows(ws, rows):
    for row in rows:
        ws.append(row)


def style_workbook(wb):
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            max_len = max(len("" if cell.value is None else str(cell.value)) for cell in col)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 55)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top")


def build_device(label, directory):
    launch = read_launch(find_single(directory, "launch_raw", "apps_launch_", ".csv"))
    device = {
        "label": label,
        "dir": directory,
        "launch": launch,
        "dex": read_csv_map(os.path.join(directory, "dexopt_before.csv"), "pkg"),
        "apk": read_csv_map(os.path.join(directory, "apk_versions_before.csv"), "pkg"),
        "sha": {},
    }
    sha_path = os.path.join(directory, "apk_sha256_before.csv")
    if os.path.exists(sha_path):
        device["sha"] = read_csv_map(sha_path, "pkg")
    return device


def make_three_sheet(wb, devices):
    ws = wb.active
    ws.title = "三设备共同app"
    labels = [d["label"] for d in devices]
    ws.append([
        "package", "activity",
        f"{labels[0]}_avg", f"{labels[1]}_avg", f"{labels[2]}_avg",
        f"{labels[1]}-{labels[0]}_avg_delta",
        f"{labels[2]}-{labels[0]}_avg_delta",
        f"{labels[2]}-{labels[1]}_avg_delta",
        f"{labels[0]}_t1", f"{labels[1]}_t1", f"{labels[2]}_t1",
        f"{labels[1]}-{labels[0]}_t1_delta",
        f"{labels[0]}_warm_t2_t5", f"{labels[1]}_warm_t2_t5", f"{labels[2]}_warm_t2_t5",
        f"{labels[1]}-{labels[0]}_warm_delta",
        f"{labels[0]}_dexopt", f"{labels[1]}_dexopt", f"{labels[2]}_dexopt",
        f"{labels[0]}_version", f"{labels[1]}_version", f"{labels[2]}_version",
        "apk_sha256_same_all3",
    ])
    pkgs = set(devices[0]["launch"])
    for device in devices[1:]:
        pkgs &= set(device["launch"])
    common = [
        pkg for pkg in sorted(pkgs)
        if all(device["launch"][pkg].get("status") == "ok" for device in devices)
    ]
    for pkg in common:
        rows = [device["launch"][pkg] for device in devices]
        warm = [warm_avg(row) for row in rows]
        shas = [sha_state(device, pkg) for device in devices]
        ws.append([
            pkg, rows[0].get("activity", ""),
            rows[0]["avg"], rows[1]["avg"], rows[2]["avg"],
            rows[1]["avg"] - rows[0]["avg"],
            rows[2]["avg"] - rows[0]["avg"],
            rows[2]["avg"] - rows[1]["avg"],
            rows[0]["t1"], rows[1]["t1"], rows[2]["t1"],
            rows[1]["t1"] - rows[0]["t1"],
            warm[0], warm[1], warm[2],
            round(warm[1] - warm[0], 1) if warm[0] is not None and warm[1] is not None else None,
            dex_state(devices[0], pkg), dex_state(devices[1], pkg), dex_state(devices[2], pkg),
            version_state(devices[0], pkg), version_state(devices[1], pkg), version_state(devices[2], pkg),
            len(set(shas)) == 1 and bool(shas[0]),
        ])
    return common


def make_pair_sheet(wb, left, right, name):
    ws = wb.create_sheet(name[:31])
    ws.append([
        "package", "activity",
        f"{left['label']}_status", f"{right['label']}_status",
        f"{left['label']}_avg", f"{right['label']}_avg",
        f"{right['label']}-{left['label']}_avg_delta",
        f"{right['label']}_vs_{left['label']}_pct",
        f"{left['label']}_t1", f"{right['label']}_t1",
        f"{right['label']}-{left['label']}_t1_delta",
        f"{left['label']}_warm_t2_t5", f"{right['label']}_warm_t2_t5",
        f"{right['label']}-{left['label']}_warm_delta",
        f"{left['label']}_dexopt", f"{right['label']}_dexopt", "dexopt_same",
        f"{left['label']}_version", f"{right['label']}_version", "version_same",
        "apk_sha256_same",
    ])
    pkgs = set(left["launch"]) & set(right["launch"])
    common = [
        pkg for pkg in sorted(pkgs)
        if left["launch"][pkg].get("status") == "ok" and right["launch"][pkg].get("status") == "ok"
    ]
    for pkg in common:
        lrow = left["launch"][pkg]
        rrow = right["launch"][pkg]
        lwarm = warm_avg(lrow)
        rwarm = warm_avg(rrow)
        avg_delta = rrow["avg"] - lrow["avg"]
        ws.append([
            pkg, lrow.get("activity", ""),
            lrow.get("status"), rrow.get("status"),
            lrow["avg"], rrow["avg"],
            avg_delta,
            round(avg_delta * 100 / lrow["avg"], 1) if lrow["avg"] else None,
            lrow["t1"], rrow["t1"],
            rrow["t1"] - lrow["t1"],
            lwarm, rwarm,
            round(rwarm - lwarm, 1) if lwarm is not None and rwarm is not None else None,
            dex_state(left, pkg), dex_state(right, pkg), dex_state(left, pkg) == dex_state(right, pkg),
            version_state(left, pkg), version_state(right, pkg), version_state(left, pkg) == version_state(right, pkg),
            sha_state(left, pkg) == sha_state(right, pkg) and bool(sha_state(left, pkg)),
        ])
    return common


def make_summary(wb, devices, three_common, pair_common):
    ws = wb.create_sheet("summary")
    labels = [d["label"] for d in devices]
    rows = [
        ["three_device_common_ok_count", len(three_common)],
        ["pair_common_ok_count", len(pair_common)],
    ]
    if len(devices) >= 2:
        left, right = devices[0], devices[1]
        for metric in ("avg", "t1"):
            for device in (left, right):
                vals = [device["launch"][pkg][metric] for pkg in pair_common]
                rows.append([f"pair_common_{device['label']}_{metric}_mean", mean(vals)])
                rows.append([f"pair_common_{device['label']}_{metric}_median", median(vals)])
        deltas = [right["launch"][pkg]["avg"] - left["launch"][pkg]["avg"] for pkg in pair_common]
        rows.append([f"{right['label']}-{left['label']}_avg_mean_delta", mean(deltas)])
        rows.append([f"{right['label']}-{left['label']}_avg_median_delta", median(deltas)])
    if three_common:
        for metric in ("avg", "t1"):
            for device in devices:
                vals = [device["launch"][pkg][metric] for pkg in three_common]
                rows.append([f"three_common_{device['label']}_{metric}_mean", mean(vals)])
                rows.append([f"three_common_{device['label']}_{metric}_median", median(vals)])
    rows.append(["device_order", " -> ".join(labels)])
    add_rows(ws, rows)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--device",
        action="append",
        required=True,
        help="LABEL=/path/to/evidence_dir. Provide 2 or 3 devices.",
    )
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    devices = []
    for item in args.device:
        if "=" not in item:
            parser.error("--device must be LABEL=/path")
        label, directory = item.split("=", 1)
        devices.append(build_device(label, directory))
    if len(devices) not in (2, 3):
        parser.error("provide 2 or 3 --device entries")

    wb = Workbook()
    if len(devices) == 3:
        three_common = make_three_sheet(wb, devices)
        pair_common = make_pair_sheet(wb, devices[0], devices[1], f"{devices[1]['label']}_{devices[0]['label']}共同app")
    else:
        wb.active.title = "placeholder"
        wb.remove(wb.active)
        three_common = []
        pair_common = make_pair_sheet(wb, devices[0], devices[1], f"{devices[1]['label']}_{devices[0]['label']}共同app")
    make_summary(wb, devices, three_common, pair_common)
    style_workbook(wb)
    wb.save(args.out)
    print(args.out)


if __name__ == "__main__":
    main()
